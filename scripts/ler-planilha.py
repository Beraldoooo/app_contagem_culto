"""
Lê um mês da planilha do Leonardo e confere contra as mensagens do grupo.

Uso (precisa de openpyxl):
  python3 -m venv /tmp/v && /tmp/v/bin/pip install openpyxl
  /tmp/v/bin/python scripts/ler-planilha.py "Relatório de contagem 2026.xlsx" set 2026 saida.json

Importa as PEÇAS (adultos, cada ministério), nunca as células de TOTAL, que
têm erro de fórmula. Imprime, culto a culto, se o que o app vai calcular bate
com o painel da mensagem detalhada da própria aba, e grava os dados em JSON.
Foi usado para gerar supabase/importar-setembro-2026.sql (21 de 21 batendo).

Formatos antigos que ele já entende:
  - "Crianças" da mensagem incluindo o Pads (até meados de set/2026);
  - Volts como "Equipe adolescentes Volts" / "Volts adolescentes".
Quando um mês novo não bater, olhar a aba antes de importar: pode ser mais um
formato antigo, e não erro de contagem.
"""
import json, re, sys, unicodedata
import openpyxl

ARQ, PREFIXO_MES, ANO = sys.argv[1], sys.argv[2], sys.argv[3]  # ex.: "set", "2026"

def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ]", " ", s).strip()

ITENS = {  # rótulo na planilha -> (codigo, grupo, soma_em)
    "estacionamento": ("estacionamento", "fora", None),
    "bercario": ("bercario", "fora", "criancas"),
    "maternal": ("maternal", "fora", "criancas"),
    "kids": ("kids", "fora", "criancas"),
    "pads": ("pads", "fora", "pads"),
    "adolescentes": ("volts", "fora", "volts"),
    "louvor": ("louvor", "dentro", None),
    "projecao": ("projecao", "dentro", None),
    "audio": ("audio", "dentro", None),
    "iluminacao": ("iluminacao", "dentro", None),
    "midias": ("midias", "dentro", None),
    "hospitalidade": ("hospitalidade", "dentro", None),
    "dizimos": ("dizimos", "dentro", None),
    "livraria": ("livraria", "dentro", None),
}
COM_PUBLICO = {"bercario", "maternal", "kids", "pads", "volts"}
GRUPO = {c: g for c, g, _ in ITENS.values()} | {"criancas_no_culto": "fora", "pastores": "pastoral"}
SOMA = {c: s for c, _, s in ITENS.values()} | {"criancas_no_culto": "criancas", "pastores": None}
MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8,
         "set": 9, "out": 10, "nov": 11, "dez": 12}
NOME_CULTO = {("dom", "08:00"): "Culto Domingo 8h", ("dom", "10:00"): "Culto Domingo 10h",
              ("dom", "18:00"): "Culto Domingo 18h", ("qua", "14:30"): "Culto da Esperança",
              ("qui", "20:00"): "Culto Quinta 20h"}
HORARIO_DIA = {"qua": "14:30", "qui": "20:00"}

problemas = []

def numero(v, onde):
    """Número inteiro, ou None se vazio. Texto no lugar de número vira problema."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        if v != int(v) or v < 0:
            problemas.append(f"{onde}: valor estranho {v!r}")
        return int(v)
    s = str(v).strip()
    if re.fullmatch(r"\d+", s):
        return int(s)
    problemas.append(f"{onde}: texto no lugar de número: {s!r}")
    return None

wb = openpyxl.load_workbook(ARQ, read_only=True, data_only=True)
cultos, conferencia = [], []

for aba in wb.sheetnames:
    m = re.match(r"\s*(dom|qua|qui)\s*0?(\d{1,2})\s*([a-z]{3})", aba, re.I)
    if not m or m.group(3).lower() != PREFIXO_MES:
        continue
    dia_sem, dia, mes = m.group(1).lower(), int(m.group(2)), MESES[m.group(3).lower()]
    data = f"{ANO}-{mes:02d}-{dia:02d}"
    rows = [list(r) for r in wb[aba].iter_rows(values_only=True)]
    cel = lambda r, c: rows[r][c] if r < len(rows) and c < len(rows[r]) else None
    txt = lambda r, c: str(cel(r, c) or "").strip()

    # Blocos de lançamento: "Adultos:" com "preencher" duas colunas à direita.
    ancoras = [(r, c) for r in range(len(rows)) for c in range(len(rows[r]))
               if norm(txt(r, c)) == "adultos" and norm(txt(r, c + 2)) == "preencher"]
    for r0, x in ancoras:
        horario = None
        for k in range(r0, max(-1, r0 - 5), -1):
            mm = re.match(r"culto das (\d+)h", norm(txt(k, x)))
            if mm:
                horario = f"{int(mm.group(1)):02d}:00"
        horario = horario or HORARIO_DIA.get(dia_sem)
        nome = NOME_CULTO.get((dia_sem, horario))
        if not nome:
            problemas.append(f"{aba}: culto não reconhecido em {r0+1}/{x+1} ({horario})")
            continue
        fim = min([r for r, c in ancoras if c == x and r > r0] + [len(rows)])
        onde = f"{aba} {horario}"
        culto = {"data": data, "horario": horario, "nome": nome,
                 "adultos": numero(cel(r0, x + 1), f"{onde} adultos"), "itens": {}}
        for r in range(r0 + 1, fim):
            rot = norm(txt(r, x))
            if rot in ITENS:
                cod = ITENS[rot][0]
                s = numero(cel(r, x + 1), f"{onde} {rot} servindo")
                p = numero(cel(r, x + 4), f"{onde} {rot} crianças") if cod in COM_PUBLICO else None
                if s is not None or p is not None:
                    culto["itens"][cod] = [s, p]
                if rot == "estacionamento":  # ao lado do estacionamento fica "crianças no culto"
                    nc = numero(cel(r, x + 4), f"{onde} crianças no culto")
                    if nc is not None:
                        culto["itens"]["criancas_no_culto"] = [None, nc]
            # Layout antigo (até set/2026): o Volts tinha linha própria com outro
            # nome — "Equipe adolescentes Volts" para quem serviu e, uma linha
            # acima, "Volts adolescentes" com o número uma coluna mais à direita.
            if rot.startswith("equipe adolescentes"):
                s = numero(cel(r, x + 1), f"{onde} volts servindo")
                if s is not None:
                    culto["itens"].setdefault("volts", [None, None])[0] = s
            if "volts" in norm(txt(r, x + 3)):
                for off in (4, 5):
                    v = cel(r, x + off)
                    if isinstance(v, (int, float)):
                        culto["itens"].setdefault("volts", [None, None])[1] = int(v)
                        break
            if norm(txt(r, x + 3)) == "pastores":
                pv = numero(cel(r, x + 4), f"{onde} pastores")
                if pv is not None:
                    culto["itens"]["pastores"] = [pv, None]
        cultos.append(culto)

    # Painel da mensagem detalhada, para conferir.
    for r in range(len(rows)):
        for c in range(len(rows[r])):
            if norm(txt(r, c)).startswith("frequencia detalhada"):
                atual = HORARIO_DIA.get(dia_sem)
                painel = {}
                for k in range(r + 1, min(len(rows), r + 45)):
                    rot = norm(txt(k, c))
                    mm = re.match(r"culto das (\d+)h", rot)
                    if mm:
                        atual = f"{int(mm.group(1)):02d}:00"
                    elif rot.startswith("total ") or rot.startswith("total geral"):
                        continue  # totais do dia: não são de um culto
                    elif rot and atual:
                        v = cel(k, c + 1)
                        if isinstance(v, (int, float)):
                            painel.setdefault(atual, {})[rot] = int(v)
                for h, campos in painel.items():
                    conferencia.append({"data": data, "horario": h, "aba": aba, "painel": campos})

def resumir(c):
    it = c["itens"]
    def soma_pub(alvo):
        vals = [p for cod, (s, p) in it.items() if SOMA.get(cod) == alvo and p is not None]
        return sum(vals) if vals else None
    serv = lambda g: sum(s or 0 for cod, (s, p) in it.items() if GRUPO[cod] == g)
    cr, pa, vo = soma_pub("criancas"), soma_pub("pads"), soma_pub("volts")
    simples = (c["adultos"] or 0) + (cr or 0) + (pa or 0) + (vo or 0)
    return {"adultos": c["adultos"], "kids": cr, "pads": pa, "adolescentes": vo,
            "servindo dentro": serv("dentro"), "servindo fora": serv("fora"),
            "equipe pastoral": serv("pastoral"), "simples": simples,
            "total": simples + serv("dentro") + serv("fora") + serv("pastoral")}

def campos_painel(p):
    alias = {"criancas": "kids", "kids": "kids", "pads": "pads", "adolescentes": "adolescentes",
             "adultos": "adultos", "servindo dentro": "servindo dentro", "servindo fora": "servindo fora",
             "equipe pastoral": "equipe pastoral", "total": "total"}
    return {alias[k]: v for k, v in p.items() if k in alias}

print(f"{len(cultos)} cultos lidos\n")
divergencias = 0
for c in sorted(cultos, key=lambda c: (c["data"], c["horario"])):
    r = resumir(c)
    conf = next((x for x in conferencia if x["data"] == c["data"] and x["horario"] == c["horario"]), None)
    dif = []
    if conf:
        painel = campos_painel(conf["painel"])
        # Até meados de set/2026 a mensagem não tinha linha de Pads: o
        # "Crianças" dela já incluía o Pads.
        if "pads" not in painel and r["pads"] is not None:
            r = {**r, "kids": (r["kids"] or 0) + r["pads"]}
        for k, v in painel.items():
            meu = r.get(k)
            if (meu or 0) != v:
                dif.append(f"{k}: app {meu} × painel {v}")
    else:
        dif.append("sem painel para conferir")
    divergencias += bool(dif)
    marca = "✓" if not dif else "≠"
    print(f"{marca} {c['data']} {c['horario']} {c['nome']:20} adultos {str(r['adultos']):>4} "
          f"simples {r['simples']:>4} detalhado {r['total']:>5}" + ("" if not dif else "   ← " + "; ".join(dif)))

print(f"\n{divergencias} cultos com divergência")
print("\nProblemas de leitura:" if problemas else "\nNenhum problema de leitura.")
for p in problemas:
    print("  -", p)

json.dump(cultos, open(sys.argv[4], "w"), ensure_ascii=False, indent=1)
